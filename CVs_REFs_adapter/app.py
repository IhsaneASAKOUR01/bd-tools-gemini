# Combined adapter: references + resumes
import os
import tempfile
from io import BytesIO
from pathlib import Path

import streamlit as st

from .docx_adapter import get_sections_from_docx, save_adapted_docx
from .gpt_logic import adapt_all_sections
from .resume_adapter import (
    adapt_resume_sections,
    extract_relevant_sections_from_resume,
    save_filtered_resume,
)


@st.cache_data(show_spinner=False)
def cached_adaptation(sections_list, ao_title):
    return adapt_all_sections(sections_list, ao_title)


@st.cache_data(show_spinner=False)
def cached_resume_adaptation(raw_sections, ao_title):
    return adapt_resume_sections(raw_sections, ao_title)


if "ref_result" not in st.session_state:
    st.session_state["ref_result"] = None

if "resume_results" not in st.session_state:
    st.session_state["resume_results"] = []


def run_app():
    title_col, action_col = st.columns([5.3, 1.3], vertical_alignment="bottom")
    with title_col:
        ao_title = st.text_input(
            "Tender / AO title",
            placeholder="Paste the tender or AO title",
        )
    with action_col:
        run_now = st.button(
            "Generate",
            key="submit_refs_cvs_btn",
            type="primary",
        )

    ref_col, cv_col = st.columns(2, gap="medium")
    with ref_col:
        uploaded_ref = st.file_uploader(
            "Project reference",
            type=["docx"],
            key="ref",
        )
    with cv_col:
        uploaded_resumes = st.file_uploader(
            "Consultant CVs",
            type=["docx"],
            accept_multiple_files=True,
            key="adapter_resumes",
        )

    if not uploaded_ref:
        st.session_state["ref_result"] = None
    if not uploaded_resumes:
        st.session_state["resume_results"] = []

    if run_now and not ao_title:
        st.warning("Add the tender / AO title first.")

    if run_now and ao_title:
        st.session_state["resume_results"] = []
        st.session_state["ref_result"] = None

        if uploaded_ref:
            with st.spinner("Tailoring reference..."):
                with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                    tmp.write(uploaded_ref.read())
                    input_path = tmp.name

                sections_list = get_sections_from_docx(input_path)
                if not sections_list:
                    st.error("No sections could be extracted from the reference file.")
                else:
                    adapted = cached_adaptation(sections_list, ao_title)
                    original_name = Path(uploaded_ref.name).stem
                    output_name = original_name + "_adapted.docx"
                    output_path = os.path.join(tempfile.gettempdir(), output_name)
                    save_adapted_docx(input_path, output_path, adapted, ao_title)

                    try:
                        from docx import Document
                        from deep_translator import GoogleTranslator
                        from langdetect import detect

                        def translate_paragraph(para, target_lang="en"):
                            full_text = para.text.strip()
                            if full_text:
                                try:
                                    translated = GoogleTranslator(source="auto", target=target_lang).translate(full_text)
                                    for run in para.runs:
                                        run.text = ""
                                    para.runs[0].text = translated
                                except Exception as e:
                                    print("[Final Translation Error]", e)

                        ao_lang = detect(ao_title)
                        doc = Document(output_path)

                        for para in doc.paragraphs:
                            translate_paragraph(para, ao_lang)

                        for table in doc.tables:
                            for row in table.rows:
                                for cell in row.cells:
                                    for para in cell.paragraphs:
                                        translate_paragraph(para, ao_lang)

                        doc.save(output_path)
                    except Exception as e:
                        print("[Final Translation Skipped]", e)

                    with open(output_path, "rb") as f:
                        st.session_state["ref_result"] = {
                            "data": f.read(),
                            "name": output_name,
                            "original": uploaded_ref.name,
                        }

        if uploaded_resumes:
            for uploaded_resume in uploaded_resumes:
                with st.spinner(f"Tailoring {uploaded_resume.name}..."):
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                        tmp.write(uploaded_resume.read())
                        resume_path = tmp.name

                    raw_sections = extract_relevant_sections_from_resume(resume_path)
                    adapted_resume = cached_resume_adaptation(raw_sections, ao_title)

                    original_name = Path(uploaded_resume.name).stem
                    output_name = original_name + "_adapted.docx"
                    output_buffer = BytesIO()
                    save_filtered_resume(resume_path, output_buffer, adapted_resume)
                    output_buffer.seek(0)
                    st.session_state["resume_results"].append(
                        {
                            "data": output_buffer.getvalue(),
                            "name": output_name,
                            "original": uploaded_resume.name,
                        }
                    )

    if st.session_state["ref_result"] or st.session_state["resume_results"]:
        st.caption("Generated files")
        downloads = []
        if st.session_state["ref_result"]:
            downloads.append(("Reference", st.session_state["ref_result"]))
        downloads.extend(("CV", result) for result in st.session_state["resume_results"])

        for kind, result in downloads:
            name_col, download_col = st.columns([5.3, 1.3], vertical_alignment="center")
            with name_col:
                st.markdown(f"**{result['original']}** · {kind}")
            with download_col:
                st.download_button(
                    "Download",
                    data=result["data"],
                    file_name=result["name"],
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
