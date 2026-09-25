from pathlib import Path
import tempfile

import streamlit as st
from docx import Document
from deep_translator import GoogleTranslator

from .template_filler import fill_reference_table, fill_template_with_debug
from .utils import load_report_text


def run_app():
    if "field_values" not in st.session_state:
        st.session_state["field_values"] = None
    if "output_path_fr" not in st.session_state:
        st.session_state["output_path_fr"] = None
    if "output_path_en" not in st.session_state:
        st.session_state["output_path_en"] = None

    input_pane, output_pane = st.columns([1.12, 0.88], gap="large")

    with input_pane:
        st.markdown("**Input**")
        uploaded_report = st.file_uploader(
            "Source report",
            type=["docx", "pdf", "pptx", "txt"],
            key="ref_creator_report",
        )
        submit = st.button(
            "Create reference",
            key="ref_creator_submit",
            type="primary",
        )

    if uploaded_report and "last_uploaded" in st.session_state:
        if uploaded_report.name != st.session_state["last_uploaded"]:
            st.session_state["generated"] = False

    st.session_state["last_uploaded"] = uploaded_report.name if uploaded_report else None

    TEMPLATE_PATH = "REF_creater/Référence Template.docx"

    def translate_docx(input_path, output_path, target_lang="en"):
        doc = Document(input_path)

        def translate_paragraph(para):
            text = para.text.strip()
            if text:
                try:
                    translated = GoogleTranslator(source="auto", target=target_lang).translate(text)
                    if para.runs:
                        para.runs[0].text = translated
                        for i in range(1, len(para.runs)):
                            para.runs[i].text = ""
                except Exception as e:
                    print("[Translation Error]", e)

        for para in doc.paragraphs:
            translate_paragraph(para)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        translate_paragraph(para)

        doc.save(output_path)

    if submit and not uploaded_report:
        st.warning("Upload a project report first.")

    if submit and uploaded_report:
        report_text = load_report_text(uploaded_report)

        with st.spinner("Extracting project information..."):
            st.session_state["field_values"] = fill_template_with_debug(TEMPLATE_PATH, report_text)

        with st.spinner("Creating French version..."):
            filled_doc = fill_reference_table(TEMPLATE_PATH, st.session_state["field_values"])

            mission_name = st.session_state["field_values"].get("Nom de la mission", "output").strip()
            mission_name_safe = "".join(c for c in mission_name if c.isalnum() or c in (" ", "_", "-")).rstrip()

            fr_name = f"{mission_name_safe}_ref_VF.docx"
            path_fr = Path(tempfile.gettempdir()) / fr_name
            filled_doc.save(path_fr)
            st.session_state["output_path_fr"] = path_fr

        with st.spinner("Creating English version..."):
            en_name = f"{mission_name_safe}_ref_VA.docx"
            path_en = Path(tempfile.gettempdir()) / en_name
            translate_docx(path_fr, path_en, target_lang="en")
            st.session_state["output_path_en"] = path_en

        st.session_state["generated"] = True

    with output_pane:
        with st.container(border=True):
            st.markdown('<div class="output-heading">Output</div>', unsafe_allow_html=True)

            paths = []
            if st.session_state["output_path_fr"]:
                paths.append(("French", st.session_state["output_path_fr"]))
            if st.session_state["output_path_en"]:
                paths.append(("English", st.session_state["output_path_en"]))

            if not paths:
                st.markdown(
                    '<div class="output-empty">The French and English reference files will appear here.</div>',
                    unsafe_allow_html=True,
                )
            else:
                for language, path in paths:
                    name_col, dl_col = st.columns([3.7, 1.0], vertical_alignment="center")
                    with name_col:
                        st.markdown(f"**{language} reference**")
                        st.caption(path.name)
                    with dl_col:
                        with open(path, "rb") as f:
                            st.download_button(
                                "Download",
                                data=f,
                                file_name=path.name,
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            )
